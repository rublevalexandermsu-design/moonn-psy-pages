from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    Flowable,
    Image as PdfImage,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "assets" / "teen-psychology-camp-2026"
DOC_DIR = ROOT / "docs" / "teen-psychology-camp-2026"
DOWNLOADS = Path.home() / "Downloads"

PUBLIC_PAGE_URL = "https://мунн.рф/podrostkovyy-lager-psihologiya"
PAYMENT_URL = f"{PUBLIC_PAGE_URL}?pay=teen-camp-2026"
TELEGRAM_URL = "https://t.me/moonn_official"

PROGRAM_PDF = ASSET_DIR / "teen-psychology-camp-tatyana-moonn-program-2026.pdf"
POSTER_PDF = ASSET_DIR / "teen-psychology-camp-tatyana-moonn-poster-2026.pdf"
POSTER_JPG = ASSET_DIR / "teen-psychology-camp-tatyana-moonn-poster-2026.jpg"
CALL_BRIEF = DOC_DIR / "teen-psychology-camp-call-brief-2026.docx"

FONT_REGULAR = "TeenCampArial"
FONT_BOLD = "TeenCampArialBold"


def register_fonts() -> None:
    fonts_dir = Path("C:/Windows/Fonts")
    regular = fonts_dir / "arial.ttf"
    bold = fonts_dir / "arialbd.ttf"
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont(FONT_REGULAR, str(regular)))
        pdfmetrics.registerFont(TTFont(FONT_BOLD, str(bold)))


def button(canvas: Canvas, x: float, y: float, w: float, h: float, label: str, url: str, fill: colors.Color) -> None:
    canvas.saveState()
    canvas.setFillColor(fill)
    canvas.roundRect(x, y, w, h, 8, stroke=0, fill=1)
    canvas.setFillColor(colors.white)
    canvas.setFont(FONT_BOLD, 13)
    canvas.drawCentredString(x + w / 2, y + h / 2 - 4, label)
    canvas.linkURL(url, (x, y, x + w, y + h), relative=0)
    canvas.restoreState()


class LinkButton(Flowable):
    def __init__(self, label: str, url: str, width: float, height: float, fill: colors.Color):
        super().__init__()
        self.label = label
        self.url = url
        self.width = width
        self.height = height
        self.fill = fill

    def draw(self) -> None:
        button(self.canv, 0, 0, self.width, self.height, self.label, self.url, self.fill)


def build_program_pdf() -> None:
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "TitleRu",
        parent=styles["Title"],
        fontName=FONT_BOLD,
        fontSize=25,
        leading=29,
        textColor=colors.HexColor("#16215c"),
        spaceAfter=8,
    )
    h2 = ParagraphStyle(
        "HeadingRu",
        parent=styles["Heading2"],
        fontName=FONT_BOLD,
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#2f3d99"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body = ParagraphStyle(
        "BodyRu",
        parent=styles["BodyText"],
        fontName=FONT_REGULAR,
        fontSize=10.7,
        leading=14.2,
        textColor=colors.HexColor("#32384d"),
        spaceAfter=6,
    )
    small = ParagraphStyle("SmallRu", parent=body, fontSize=9.5, leading=12)
    cell = ParagraphStyle("CellRu", parent=body, fontSize=8.7, leading=11)
    cell_bold = ParagraphStyle("CellBoldRu", parent=cell, fontName=FONT_BOLD)

    doc = SimpleDocTemplate(
        str(PROGRAM_PDF),
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=15 * mm,
        bottomMargin=16 * mm,
        title="Психология без скуки - программа подросткового интенсива",
    )
    story = [
        Paragraph("Психология без скуки: подростковый интенсив", title),
        Paragraph("6-10 июля 2026 · Москва, Цветной бульвар · 10:00-14:00 · группа 10-12 подростков", body),
        Paragraph("Ранняя оплата до 15 июня: <b>40 000 ₽</b> вместо <strike>50 000 ₽</strike>. Формат: психологические и развивающие занятия без централизованного питания.", body),
    ]

    cta_table = Table(
        [[
            LinkButton("Оплатить участие", PAYMENT_URL, 50 * mm, 12 * mm, colors.HexColor("#3149c9")),
            LinkButton("Страница интенсива", PUBLIC_PAGE_URL, 50 * mm, 12 * mm, colors.HexColor("#7b4be0")),
            LinkButton("Вопрос в Telegram", TELEGRAM_URL, 50 * mm, 12 * mm, colors.HexColor("#12a896")),
        ]],
        colWidths=[52 * mm, 52 * mm, 52 * mm],
        rowHeights=[13 * mm],
    )
    cta_table.setStyle(
        TableStyle(
            [
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0, colors.white),
            ]
        )
    )
    story += [Spacer(1, 5 * mm), cta_table, Spacer(1, 4 * mm)]

    story += [
        Paragraph("Для кого", h2),
        Paragraph("Для подростков, которым важно спокойнее общаться, увереннее говорить о себе, управлять тревогой, пробовать нейросети для идей и понять первые шаги в профессии психолога.", body),
        Paragraph("Расписание дня", h2),
    ]

    rows = [[Paragraph("Время", cell_bold), Paragraph("Что происходит", cell_bold)]]
    schedule = [
        ("10:00-10:20", "Сбор, настрой на день и правила безопасного общения."),
        ("10:20-11:20", "Тема дня: короткий разбор, упражнения и личная карта наблюдений."),
        ("11:20-11:35", "Пауза, вода, восстановление внимания."),
        ("11:35-12:45", "Практика: общение, эмоции, уверенность, ИИ или профессия психолога."),
        ("12:45-14:00", "Закрепление, обратная связь и личный план на день."),
    ]
    rows += [[Paragraph(t, cell_bold), Paragraph(text, cell)] for t, text in schedule]
    table = Table(rows, colWidths=[34 * mm, 128 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1ff")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d7dcee")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(table)

    story += [Paragraph("Пять дней", h2)]
    days = [
        ("День 1", "Знакомство, доверие и безопасная атмосфера: правила группы, упражнения на контакт, первая карта целей."),
        ("День 2", "Эмоции и тревога: как замечать напряжение, снижать стресс и говорить о сложных состояниях."),
        ("День 3", "Общение и уверенность: практика диалогов, границы, просьбы, обратная связь."),
        ("День 4", "ИИ для идей и проектов: как использовать нейросети экологично, не заменяя собственное мышление."),
        ("День 5", "Первые шаги молодого психолога и личный план: что делает психолог, какие качества важны, как продолжить развитие."),
    ]
    day_rows = [[Paragraph(d, cell_bold), Paragraph(text, cell)] for d, text in days]
    day_table = Table(day_rows, colWidths=[25 * mm, 137 * mm])
    day_table.setStyle(table._cellStyles and TableStyle([("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d7dcee")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(day_table)

    story += [
        Paragraph("Календарь следующих потоков", h2),
        Paragraph("24-28 августа — «Разобраться в себе»; 19-20 и 26-27 сентября — «Уверенность в себе»; 26-30 октября — «Личная эффективность»; 2-6 ноября — «Выбор профессии». Финальные даты подтверждаются перед запуском потока.", body),
        Spacer(1, 5 * mm),
        Paragraph("Нажмите кнопку ниже, чтобы сразу открыть оплату ранней стоимости или задать вопрос организатору.", small),
    ]

    def footer_canvas(canvas: Canvas, document) -> None:
        canvas.saveState()
        w, _ = A4
        button(canvas, 20 * mm, 10 * mm, 54 * mm, 10 * mm, "Оплатить участие", PAYMENT_URL, colors.HexColor("#3149c9"))
        button(canvas, 78 * mm, 10 * mm, 54 * mm, 10 * mm, "Страница интенсива", PUBLIC_PAGE_URL, colors.HexColor("#7b4be0"))
        button(canvas, 136 * mm, 10 * mm, 54 * mm, 10 * mm, "Вопрос в Telegram", TELEGRAM_URL, colors.HexColor("#12a896"))
        canvas.setFont(FONT_REGULAR, 8)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawRightString(w - 16 * mm, 6 * mm, "Moonn · подростковый интенсив · 2026")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer_canvas, onLaterPages=footer_canvas)


def build_poster_pdf() -> None:
    img = Image.open(POSTER_JPG)
    page_w, page_h = A4
    canvas = Canvas(str(POSTER_PDF), pagesize=A4)
    margin = 10 * mm
    target_w = page_w - 2 * margin
    target_h = page_h - 47 * mm
    ratio = min(target_w / img.width, target_h / img.height)
    draw_w = img.width * ratio
    draw_h = img.height * ratio
    x = (page_w - draw_w) / 2
    y = 37 * mm
    canvas.drawImage(str(POSTER_JPG), x, y, draw_w, draw_h, preserveAspectRatio=True, mask="auto")
    canvas.setFillColor(colors.HexColor("#f4efff"))
    canvas.roundRect(10 * mm, 9 * mm, page_w - 20 * mm, 23 * mm, 9, stroke=0, fill=1)
    canvas.setFont(FONT_BOLD, 12)
    canvas.setFillColor(colors.HexColor("#16215c"))
    canvas.drawString(16 * mm, 25 * mm, "Ранняя оплата до 15 июня: 40 000 ₽ вместо 50 000 ₽")
    button(canvas, 16 * mm, 12 * mm, 53 * mm, 10 * mm, "Оплатить участие", PAYMENT_URL, colors.HexColor("#3149c9"))
    button(canvas, 73 * mm, 12 * mm, 53 * mm, 10 * mm, "Страница интенсива", PUBLIC_PAGE_URL, colors.HexColor("#7b4be0"))
    button(canvas, 130 * mm, 12 * mm, 53 * mm, 10 * mm, "Вопрос в Telegram", TELEGRAM_URL, colors.HexColor("#12a896"))
    canvas.save()


def add_hyperlink(paragraph, text: str, url: str, color: str = "3149C9", bold: bool = False) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    if bold:
        r_pr.append(OxmlElement("w:b"))
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9.5)


def build_call_brief_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_p.add_run("Памятка для созвона: подростковый интенсив «Психология без скуки»")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(22, 33, 92)

    lead = doc.add_paragraph()
    lead.add_run("Коротко: ").bold = True
    lead.add_run("психологический интенсив Татьяны Мунн для подростков: уверенность, общение, эмоции, ИИ для идей и первые шаги молодого психолога. 6-10 июля, Цветной бульвар, 10:00-14:00, группа 10-12 человек. Формат подается как психологические и развивающие занятия без централизованного питания.")

    p = doc.add_paragraph()
    p.add_run("Главная страница интенсива: ").bold = True
    add_hyperlink(p, PUBLIC_PAGE_URL, PUBLIC_PAGE_URL)
    p = doc.add_paragraph()
    p.add_run("Оплатить участие: ").bold = True
    add_hyperlink(p, "открыть оплату 40 000 ₽ до 15 июня", PAYMENT_URL, bold=True)
    p = doc.add_paragraph()
    p.add_run("Вопросы в Telegram: ").bold = True
    add_hyperlink(p, TELEGRAM_URL, TELEGRAM_URL)

    for heading, items in [
        ("Кому подходит", [
            "подростку трудно общаться или проявляться в группе;",
            "есть тревога перед учебой, экзаменами или новыми людьми;",
            "родители хотят практичный формат без тяжелой лекционности;",
            "ребенку интересны психология, нейросети, самопознание или выбор профессии.",
        ]),
        ("Что сказать в первые 30 секунд", [
            "Это камерный психологический интенсив: серия развивающих занятий с ведущим специалистом.",
            "Фокус: уверенность, общение, эмоции, ИИ для идей, личный план и знакомство с профессией психолога.",
            "Татьяна Мунн ведет программу лично, группа 10-12 подростков.",
        ]),
        ("Цена и оплата", [
            "ранняя стоимость до 15 июня — 40 000 ₽;",
            "стандартная стоимость — 50 000 ₽;",
            "оплата открывается по ссылке выше на странице интенсива;",
            "после оплаты нужно написать в Telegram возраст подростка и ожидания.",
        ]),
    ]:
        doc.add_heading(heading, level=1)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Расписание дня", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_text(table.rows[0].cells[0], "Время", True)
    set_cell_text(table.rows[0].cells[1], "Что происходит", True)
    for time, text in [
        ("10:00-10:20", "сбор, настрой на день и правила безопасного общения"),
        ("10:20-11:20", "тема дня, упражнения и личная карта наблюдений"),
        ("11:20-11:35", "пауза, вода, восстановление внимания"),
        ("11:35-12:45", "практика общения, эмоций, уверенности, ИИ или профессии"),
        ("12:45-14:00", "рефлексия, обратная связь, личный план"),
    ]:
        row = table.add_row().cells
        set_cell_text(row[0], time, True)
        set_cell_text(row[1], text)

    doc.add_heading("Календарь потоков", level=1)
    for item in [
        "6-10 июля — базовый поток «Психология без скуки».",
        "24-28 августа — «Разобраться в себе».",
        "19-20 и 26-27 сентября — «Уверенность в себе».",
        "26-30 октября — «Личная эффективность».",
        "2-6 ноября — «Выбор профессии».",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Материалы", level=1)
    p = doc.add_paragraph()
    p.add_run("Программа PDF: ").bold = True
    add_hyperlink(p, "скачать с сайта", f"{PUBLIC_PAGE_URL}#materials")
    p = doc.add_paragraph()
    p.add_run("Постер PDF: ").bold = True
    add_hyperlink(p, "скачать с сайта", f"{PUBLIC_PAGE_URL}#materials")
    p = doc.add_paragraph()
    p.add_run("CTA для отправки клиенту: ").bold = True
    add_hyperlink(p, "Оплатить участие в интенсиве", PAYMENT_URL, bold=True)

    doc.add_heading("Частые вопросы", level=1)
    qa = [
        ("Нужно ли родителям присутствовать?", "Занятия проходят в группе с ведущим специалистом; организационные условия подтверждаются при записи."),
        ("Есть ли питание?", "Централизованное питание не заявляется. В расписании есть короткая пауза; индивидуальные вопросы лучше уточнить до старта."),
        ("Можно ли записаться на другой поток?", "Да, есть августовская, сентябрьская, октябрьская и ноябрьская предзапись."),
        ("Что делать после оплаты?", "Написать в Telegram возраст подростка, контакт родителя и ожидания от участия."),
    ]
    for q, a in qa:
        p = doc.add_paragraph()
        p.add_run(q + " ").bold = True
        p.add_run(a)

    doc.save(str(CALL_BRIEF))


def replace_downloads() -> None:
    old_names = [
        "moonn-teen-camp-program-2026.pdf",
        "moonn-teen-camp-poster-2026.jpg",
        "moonn-teen-camp-call-brief-2026.docx",
        "Программа подросткового лагеря Татьяны Мунн 2026.pdf",
        "Постер подросткового лагеря Татьяны Мунн 2026.pdf",
        "Памятка для созвона по подростковому лагерю Татьяны Мунн 2026.docx",
    ]
    for name in old_names:
        path = DOWNLOADS / name
        if path.exists():
            try:
                path.unlink()
            except PermissionError:
                # Word/preview may keep a previous copy open; keep generating the new canonical files.
                pass
    shutil.copy2(PROGRAM_PDF, DOWNLOADS / "Программа подросткового интенсива Татьяны Мунн 2026.pdf")
    shutil.copy2(POSTER_PDF, DOWNLOADS / "Постер подросткового интенсива Татьяны Мунн 2026.pdf")
    shutil.copy2(CALL_BRIEF, DOWNLOADS / "Памятка для созвона по подростковому интенсиву Татьяны Мунн 2026.docx")


def main() -> None:
    register_fonts()
    build_program_pdf()
    build_poster_pdf()
    build_call_brief_docx()
    replace_downloads()


if __name__ == "__main__":
    main()
